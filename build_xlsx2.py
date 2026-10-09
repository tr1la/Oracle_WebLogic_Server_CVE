#!/usr/bin/env python3
"""Dựng xlsx cuối từ wls-live.csv (108 host sống).
Cột: URL (hyperlink, domain xếp đầu) · IP:PORT · Version · Note · Thành phố · Tổ chức.
Chạy: python3 build_xlsx2.py
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
BASE = os.path.dirname(os.path.abspath(__file__))
INFILE = sys.argv[1] if len(sys.argv) > 1 else "wls-live.csv"
OUTXLSX = "WLS-VN-live-vhost.xlsx" if "vhost" in INFILE else "WLS-VN-live.xlsx"

EXCLUDE_IPS = {
    "117.122.125.107",   # IP deprecated, forward về landing của hosting provider -> false positive
}

ORG_MAP = {
    "越南邮政电信集团": "VNPT Corp",
    "越南电信": "Viettel Group",
    "越南互联网络信息中心": "Vietnam Internet Network Information Center (VNNIC)",
}

rows = []
with open(os.path.join(BASE, INFILE), newline='', encoding='utf-8') as f:
    rows = list(csv.DictReader(f))
rows = [r for r in rows if (r.get('ip') or '').strip() not in EXCLUDE_IPS]   # bỏ IP false-positive
for r in rows:
    o = (r.get('org') or '').strip()
    if o in ORG_MAP: r['org'] = ORG_MAP[o]

def has_letter(u):
    host = u.split('://', 1)[-1].split('/', 1)[0].rsplit(':', 1)[0]
    return any(c.isalpha() for c in host)

def ip_key(ip):
    try: return tuple(int(x) for x in ip.split('.'))
    except Exception: return (999, 999, 999, 999)

def port_key(p):
    try: return int(p)
    except Exception: return 0

# Ưu tiên URL có domain (chữ) lên đầu; trong mỗi nhóm (domain / IP-only) gom theo IP -> PORT
rows.sort(key=lambda r: (not has_letter(r['url']), ip_key(r['ip']), port_key(r['port']), r['url'].lower()))

wb = Workbook(); ws = wb.active; ws.title = "WLS VN (live)"
headers = ["URL", "IP:PORT", "Version", "Note", "Thành phố", "Tổ chức"]
ws.append(headers)
hf = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="C0392B")
thin = Side(style="thin", color="DDDDDD"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c in ws[1]:
    c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border

dom_fill = PatternFill("solid", fgColor="FDF2E9")
open_fill = PatternFill("solid", fgColor="FADBD8")   # tô đỏ nhạt cho console-open
link_font = Font(color="0563C1", underline="single")
for r in rows:
    ipport = f"{r['ip']}:{r['port']}"
    ws.append([r['url'], ipport, r.get('version', ''), r.get('note', ''),
               r.get('city', ''), r.get('org', '')])

for i, r in enumerate(rows, start=2):
    for c in ws[i]:
        c.border = border; c.alignment = Alignment(vertical="center")
        if has_letter(r['url']): c.fill = dom_fill
    # tô nổi bật host console-open
    if 'console-open' in (r.get('note') or ''):
        ws.cell(row=i, column=4).fill = open_fill
    a = ws.cell(row=i, column=1); a.hyperlink = r['url']; a.font = link_font

for j, w in enumerate([46, 24, 14, 22, 18, 46], start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:F{len(rows)+1}"
wb.save(os.path.join(BASE, OUTXLSX))

ndom = sum(1 for r in rows if has_letter(r['url']))
nver = sum(1 for r in rows if r.get('version'))
nopen = sum(1 for r in rows if 'console-open' in (r.get('note') or ''))
nt3 = sum(1 for r in rows if 'T3-open' in (r.get('note') or ''))
print(f"Da ghi {OUTXLSX} | host: {len(rows)} | domain: {ndom} | version: {nver} | console-open: {nopen} | T3-open: {nt3}")
