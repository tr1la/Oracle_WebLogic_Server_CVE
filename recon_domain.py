#!/usr/bin/env python3
"""
recon_domain.py — IP→Domain recon pipeline

Both modes write results directly into Weblogic-VN-live-final.xlsx.

Full recon:
    python3 recon_domain.py
    4 phases: PTR+TLS → RapidDNS → crt.sh CT → httpx+nuclei

Probe mode (add your own domains, re-probe):
    python3 recon_domain.py domains.csv
    Reads CSV (IP:PORT,Domain), runs httpx+nuclei, merges into the xlsx.
    Appends new domains, overwrites httpx/WLS columns with fresh results.
"""

import json
import os
import re
import socket
import ssl
import subprocess
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from pathlib import Path
from urllib.request import urlopen, Request

import dns.resolver
import dns.reversename
import openpyxl

TIMEOUT = 5
MAX_WORKERS = 20

NUCLEI_WLS_TEMPLATES = [
    "http/technologies/weblogic-detect.yaml",
    "http/exposed-panels/weblogic-login.yaml",
    "network/detection/weblogic-t3-detect.yaml",
    "network/detection/weblogic-iiop-detect.yaml",
]

# ── PTR lookup ───────────────────────────────────────────────────────

def ptr_lookup(ip: str) -> list[str]:
    try:
        rev = dns.reversename.from_address(ip)
        answers = dns.resolver.resolve(rev, "PTR", lifetime=TIMEOUT)
        return [str(rr).rstrip(".") for rr in answers]
    except Exception:
        return []

# ── TLS cert: CN + SAN ──────────────────────────────────────────────

def tls_grab_domains(ip: str, port: int) -> list[str]:
    domains = []
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    try:
        with socket.create_connection((ip, port), timeout=TIMEOUT) as sock:
            with ctx.wrap_socket(sock, server_hostname=ip) as ssock:
                cert = ssock.getpeercert(binary_form=True)
                # decode with the ssl module
                pem_cert = ssl.DER_cert_to_PEM_cert(cert)
                # parse via openssl-like extraction from the binary cert
                decoded = _parse_cert_domains(cert)
                domains.extend(decoded)
    except Exception:
        pass
    return domains


def _parse_cert_domains(der_cert: bytes) -> list[str]:
    """Extract CN and SAN DNS names from a DER certificate."""
    domains = set()
    try:
        pem = ssl.DER_cert_to_PEM_cert(der_cert)
        with tempfile.NamedTemporaryFile(mode="w", suffix=".pem", delete=False) as f:
            f.write(pem)
            tmp = f.name
        out = subprocess.check_output(
            ["openssl", "x509", "-in", tmp, "-noout", "-text"],
            stderr=subprocess.DEVNULL, timeout=5
        ).decode(errors="replace")
        os.unlink(tmp)
        cn_m = re.search(r"Subject:.*?CN\s*=\s*([^\s/,]+)", out)
        if cn_m:
            cn = cn_m.group(1).strip()
            if not _is_ip(cn):
                domains.add(cn.lower())
        for m in re.finditer(r"DNS:([^\s,]+)", out):
            d = m.group(1).strip().lower()
            if not _is_ip(d):
                domains.add(d)
    except Exception:
        pass
    return list(domains)


def _is_ip(s: str) -> bool:
    try:
        socket.inet_aton(s)
        return True
    except OSError:
        pass
    try:
        socket.inet_pton(socket.AF_INET6, s)
        return True
    except OSError:
        return False

# ── crt.sh CLI (az7rb/crt.sh — 4 CT sources in parallel) ────────────

def crtsh_lookup(domain: str) -> list[str]:
    """Query 4 CT log sources via crt.sh CLI for subdomains."""
    try:
        proc = subprocess.run(
            ["crt.sh", "-d", domain, "-q", "-t", "20"],
            capture_output=True, text=True, timeout=60
        )
        names: set[str] = set()
        for line in proc.stdout.strip().splitlines():
            name = line.strip().lower().lstrip("*.")
            if name and "." in name and not _is_ip(name):
                names.add(name)
        return list(names)
    except (subprocess.TimeoutExpired, FileNotFoundError):
        return []


# ── RapidDNS reverse IP lookup ──────────────────────────────────────

_RAPIDDNS_RE = re.compile(r"<td>([a-zA-Z0-9._-]+\.[a-zA-Z]{2,})</td>")

def rapiddns_reverse(ip: str) -> list[str]:
    """Reverse IP lookup via rapiddns.io — returns domains whose A record points to this IP."""
    try:
        url = f"https://rapiddns.io/sameip/{ip}?full=1"
        req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
        resp = urlopen(req, timeout=15)
        html = resp.read().decode(errors="replace")
        domains: set[str] = set()
        for m in _RAPIDDNS_RE.finditer(html):
            d = m.group(1).strip().lower()
            if d and not _is_ip(d):
                domains.add(d)
        return list(domains)
    except Exception:
        return []


def _extract_base_domains(domains: set[str]) -> set[str]:
    """Extract registrable base domains (e.g. 'ssc.gov.vn' from 'ids.ssc.gov.vn')."""
    known_suffixes = {
        "gov.vn", "com.vn", "edu.vn", "org.vn", "net.vn", "ac.vn",
        "info.vn", "biz.vn", "pro.vn", "health.vn",
        "com", "net", "org", "io", "co", "vn", "art", "mobi",
    }
    bases: set[str] = set()
    for d in domains:
        parts = d.strip(".").split(".")
        for i in range(len(parts) - 1):
            suffix = ".".join(parts[i + 1:])
            if suffix in known_suffixes:
                bases.add(".".join(parts[i:]))
                break
        else:
            if len(parts) >= 2:
                bases.add(".".join(parts[-2:]))
    return bases



# ── httpx probe ──────────────────────────────────────────────────────

def _run_httpx(targets: list[str]) -> dict[str, dict]:
    """Run httpx on a list of URLs, return {url: {status, title, tech, server}}."""
    results = {}
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False) as f:
        f.write("\n".join(targets))
        target_file = f.name

    print(f"    httpx: {len(targets)} targets, timeout=6s, threads=30")
    t0 = time.time()
    try:
        proc = subprocess.run(
            ["httpx", "-l", target_file, "-silent",
             "-status-code", "-title", "-tech-detect", "-web-server",
             "-json", "-no-color",
             "-timeout", "6", "-retries", "1", "-threads", "30"],
            capture_output=True, text=True, timeout=600
        )
        elapsed = time.time() - t0
        print(f"    httpx finished in {elapsed:.1f}s (exit code {proc.returncode})")

        if proc.stderr.strip():
            for line in proc.stderr.strip().splitlines()[-5:]:
                print(f"    httpx stderr: {line}")

        parsed = 0
        errors = 0
        for line in proc.stdout.strip().splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                url = obj.get("url", obj.get("input", "")).rstrip("/")
                status = obj.get("status_code", "")
                title = obj.get("title", "")
                server = obj.get("webserver", "")
                tech = ", ".join(obj.get("tech", []))
                content_len = obj.get("content_length", "")
                resp_time = obj.get("time", "")
                tls_info = obj.get("tls", {})
                tls_cn = tls_info.get("subject_cn", "") if isinstance(tls_info, dict) else ""

                detail = f"[{status}]"
                if title:
                    detail += f" title=\"{title}\""
                if server:
                    detail += f" server={server}"
                if tech:
                    detail += f" tech=[{tech}]"
                if content_len:
                    detail += f" len={content_len}"
                if resp_time:
                    detail += f" time={resp_time}"
                if tls_cn:
                    detail += f" tls_cn={tls_cn}"
                print(f"      {url} {detail}")

                results[url] = {
                    "status": status,
                    "title": title,
                    "tech": tech,
                    "server": server,
                }
                parsed += 1
            except json.JSONDecodeError:
                errors += 1
                continue
        print(f"    httpx summary: {parsed} live, {errors} parse errors, {len(targets) - parsed} no response")

    except subprocess.TimeoutExpired:
        elapsed = time.time() - t0
        print(f"    [!] httpx process timed out after {elapsed:.0f}s (limit=600s)")
        print(f"    [!] hint: reduce targets or increase per-host timeout")
    except FileNotFoundError:
        print("    [!] httpx not found — install: go install github.com/projectdiscovery/httpx/cmd/httpx@latest")
    finally:
        os.unlink(target_file)
    return results

# ── nuclei WebLogic fingerprint ──────────────────────────────────────

def _run_nuclei(targets: list[str]) -> dict[str, list[str]]:
    """Run nuclei WLS detection templates, return {url: [matched_template_ids]}."""
    results: dict[str, list[str]] = {}
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False) as f:
        f.write("\n".join(targets))
        target_file = f.name

    template_args = []
    for t in NUCLEI_WLS_TEMPLATES:
        template_args.extend(["-t", t])

    try:
        proc = subprocess.run(
            ["nuclei", "-l", target_file, "-silent",
             "-jsonl", "-no-color",
             "-timeout", "8", "-retries", "1", "-concurrency", "20",
             *template_args],
            capture_output=True, text=True, timeout=300
        )
        for line in proc.stdout.strip().splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                url = obj.get("url", "").rstrip("/")
                tpl = obj.get("template-id", "")
                if not url:
                    continue
                if tpl and tpl not in results.get(url, []):
                    results.setdefault(url, []).append(tpl)
            except json.JSONDecodeError:
                continue
    except subprocess.TimeoutExpired:
        print("    [!] nuclei timed out")
    except FileNotFoundError:
        print("    [!] nuclei not found, skipping fingerprint")
    finally:
        os.unlink(target_file)
    return results

# ── Excel processing ─────────────────────────────────────────────────

def copy_cell_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)
        dst.alignment = copy(src.alignment)


def process(input_path: str, row_range: tuple[int, int] | None = None):
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active

    headers = [c.value for c in ws[1]]
    ip_col_idx = None
    for i, h in enumerate(headers):
        if h and "IP" in str(h).upper() and "PORT" in str(h).upper():
            ip_col_idx = i + 1  # 1-based
            break
    if ip_col_idx is None:
        print("ERROR: Could not find IP:PORT column")
        sys.exit(1)

    # find or create Domain/httpx Probe/WLS Confirm columns
    domain_col = None
    httpx_col = None
    wls_col = None
    for i, h in enumerate(headers):
        if h == "Domain":
            domain_col = i + 1
        elif h == "httpx Probe":
            httpx_col = i + 1
        elif h == "WLS Confirm":
            wls_col = i + 1

    if domain_col is None:
        insert_at = ip_col_idx + 1
        ws.insert_cols(insert_at, amount=3)
        domain_col = insert_at
        httpx_col = insert_at + 1
        wls_col = insert_at + 2
        src_hdr = ws.cell(row=1, column=ip_col_idx)
        for ci, name in [(domain_col, "Domain"), (httpx_col, "httpx Probe"), (wls_col, "WLS Confirm")]:
            hdr = ws.cell(row=1, column=ci, value=name)
            copy_cell_style(src_hdr, hdr)
        for ci, w in [(domain_col, 50), (httpx_col, 40), (wls_col, 35)]:
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w
        # re-read headers after insert
        headers = [c.value for c in ws[1]]

    # determine row range to process
    min_row = 2
    max_row = ws.max_row
    if row_range:
        min_row = max(row_range[0], 2)  # never include header
        max_row = min(row_range[1], ws.max_row)
        print(f"[*] Processing rows {min_row}–{max_row} only")

    print(f"[*] IP:PORT col={ip_col_idx}, Domain col={domain_col}, httpx col={httpx_col}, WLS col={wls_col}")
    print(f"[*] Total data rows: {ws.max_row - 1} (processing {max_row - min_row + 1})")

    # ── Collect unique IPs (from selected rows only) ─────────────────
    ip_port_map: dict[str, tuple[str, int]] = {}
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                            min_col=ip_col_idx, max_col=ip_col_idx, values_only=True):
        val = row[0]
        if not val:
            continue
        val = str(val).strip()
        if ":" in val:
            parts = val.rsplit(":", 1)
            ip, port = parts[0], int(parts[1])
        else:
            ip, port = val, 443
        ip_port_map[val] = (ip, port)

    unique_ips = {v[0] for v in ip_port_map.values()}
    print(f"[*] Unique IPs: {len(unique_ips)}")

    # ── Phase 1: PTR + TLS (parallel) ───────────────────────────────
    print("[*] Phase 1/4: PTR + TLS cert lookup (parallel)...")
    results: dict[str, set[str]] = {val: set() for val in ip_port_map}

    def _ptr_tls(ip_port_str):
        ip, port = ip_port_map[ip_port_str]
        domains = set()
        domains.update(ptr_lookup(ip))
        if port >= 443:
            domains.update(tls_grab_domains(ip, port))
        return ip_port_str, domains

    done = 0
    total = len(ip_port_map)
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        futs = {pool.submit(_ptr_tls, k): k for k in ip_port_map}
        for fut in as_completed(futs):
            ip_port_str, domains = fut.result()
            results[ip_port_str].update(domains)
            done += 1
            if done % 10 == 0 or done == total:
                print(f"    [{done}/{total}]")

    # ── Phase 2: RapidDNS reverse IP (IP → domains) ──────────────────
    print(f"[*] Phase 2/4: RapidDNS reverse IP ({len(unique_ips)} IPs)...")
    ip_to_rapid: dict[str, list[str]] = {}
    done = 0
    total_ips = len(unique_ips)
    with ThreadPoolExecutor(max_workers=10) as pool:
        futs = {pool.submit(rapiddns_reverse, ip): ip for ip in sorted(unique_ips)}
        for fut in as_completed(futs):
            ip = futs[fut]
            ip_to_rapid[ip] = fut.result()
            done += 1
            if done % 5 == 0 or done == total_ips:
                found = sum(len(v) for v in ip_to_rapid.values())
                print(f"    [{done}/{total_ips}] ({found} domains so far)")

    for val, (ip, port) in ip_port_map.items():
        results[val].update(ip_to_rapid.get(ip, []))

    # ── Phase 3: crt.sh CT lookup (from all discovered base domains) ─
    all_domains_so_far: set[str] = set()
    for doms in results.values():
        all_domains_so_far.update(doms)
    base_domains = _extract_base_domains(all_domains_so_far)

    url_col_idx = None
    for i, h in enumerate(headers):
        if h and str(h).upper() == "URL":
            url_col_idx = i + 1
            break
    if url_col_idx:
        from urllib.parse import urlparse
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                min_col=url_col_idx, max_col=url_col_idx, values_only=True):
            val = row[0]
            if not val:
                continue
            try:
                host = urlparse(str(val)).hostname or ""
                if host and not _is_ip(host):
                    base_domains.update(_extract_base_domains({host}))
            except Exception:
                pass

    base_domains.discard("localhost")
    base_domains = {d for d in base_domains if "." in d and len(d) > 3}
    print(f"[*] Phase 3/4: crt.sh CT lookup ({len(base_domains)} base domains)...")
    crtsh_all: dict[str, list[str]] = {}
    done = 0
    total_bd = len(base_domains)
    with ThreadPoolExecutor(max_workers=10) as pool:
        futs = {pool.submit(crtsh_lookup, bd): bd for bd in sorted(base_domains)}
        for fut in as_completed(futs):
            bd = futs[fut]
            crtsh_all[bd] = fut.result()
            done += 1
            if done % 5 == 0 or done == total_bd:
                found = sum(len(v) for v in crtsh_all.values())
                print(f"    [{done}/{total_bd}] ({found} names so far)")

    for val in ip_port_map:
        row_domains = results[val]
        row_bases = _extract_base_domains(row_domains)
        if url_col_idx:
            from urllib.parse import urlparse
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
                if str(row[1] or "").strip() == val and row[0]:
                    try:
                        host = urlparse(str(row[0])).hostname or ""
                        if host and not _is_ip(host):
                            row_bases.update(_extract_base_domains({host}))
                    except Exception:
                        pass
        for bd in row_bases:
            if bd in crtsh_all:
                results[val].update(crtsh_all[bd])

    # ── Phase 4: httpx probe + nuclei fingerprint ────────────────────
    # probe base domains only (not every subdomain from crt.sh)
    print("[*] Phase 4/4: httpx probe + nuclei fingerprint...")
    all_probe_targets: set[str] = set()
    for val, domains in results.items():
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        bases = _extract_base_domains({d for d in domains if "api count" not in d and "error" not in d})
        for d in bases:
            all_probe_targets.add(f"{scheme}://{d}:{port}")
            all_probe_targets.add(f"{scheme}://{d}")
        all_probe_targets.add(f"{scheme}://{ip}:{port}")

    httpx_live: dict[str, dict] = {}
    if all_probe_targets:
        httpx_live = _run_httpx(sorted(all_probe_targets))
    print(f"    httpx: {len(httpx_live)} live targets")

    # build per-row: only live domains
    httpx_per_row: dict[str, str] = {}
    live_domains_per_row: dict[str, set[str]] = {}
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        bases = _extract_base_domains({d for d in results.get(val, set()) if "api count" not in d and "error" not in d})
        live_domains = []
        live_set: set[str] = set()
        for d in sorted(bases):
            for sc in (scheme, "https" if scheme == "http" else "http"):
                key = f"{sc}://{d}:{port}"
                if key in httpx_live:
                    info = httpx_live[key]
                    live_domains.append(f"{d} [{info.get('status','')}]")
                    live_set.add(d)
                    break
                key2 = f"{sc}://{d}"
                if key2 in httpx_live:
                    info = httpx_live[key2]
                    live_domains.append(f"{d} [{info.get('status','')}]")
                    live_set.add(d)
                    break
        if not live_domains:
            key = f"{scheme}://{ip}:{port}"
            if key in httpx_live:
                info = httpx_live[key]
                live_domains.append(f"{ip} [{info.get('status','')}]")
        httpx_per_row[val] = ", ".join(live_domains) if live_domains else ""
        live_domains_per_row[val] = live_set

    # ── nuclei WebLogic fingerprint ─────────────────────────────────
    print("[*] nuclei WebLogic fingerprint...")
    nuclei_targets: list[str] = []
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        nuclei_targets.append(f"{scheme}://{ip}:{port}")

    nuclei_results = _run_nuclei(sorted(set(nuclei_targets)))
    print(f"    nuclei: {len(nuclei_results)} confirmed WebLogic")

    nuclei_per_row: dict[str, str] = {}
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443,7002,8443,9002) else "http"
        key = f"{scheme}://{ip}:{port}"
        matches = nuclei_results.get(key, [])
        nuclei_per_row[val] = ", ".join(matches) if matches else ""

    # ── Write results back into sheet ──────────────────────────────
    print("[*] Writing results to Excel...")
    found_count = 0
    wls_count = 0
    for row_idx in range(min_row, max_row + 1):
        ip_cell = ws.cell(row=row_idx, column=ip_col_idx)
        val = str(ip_cell.value).strip() if ip_cell.value else ""

        # Domain: merge all discovered base domains with existing
        all_bases = _extract_base_domains({d for d in results.get(val, set()) if "api count" not in d and "error" not in d})
        existing = str(ws.cell(row=row_idx, column=domain_col).value or "")
        existing_set = {d.strip().lower() for d in existing.split(",") if d.strip()}
        merged = existing_set | all_bases
        domain_str = ", ".join(sorted(merged)) if merged else ""
        ws.cell(row=row_idx, column=domain_col, value=domain_str)
        if domain_str:
            found_count += 1

        # httpx Probe: overwrite
        ws.cell(row=row_idx, column=httpx_col, value=httpx_per_row.get(val, ""))

        # WLS Confirm: overwrite
        wls_val = nuclei_per_row.get(val, "")
        ws.cell(row=row_idx, column=wls_col, value=wls_val)
        if wls_val:
            wls_count += 1

    wb.save(input_path)
    row_count = max_row - min_row + 1
    print(f"\n[+] Done! Saved to {input_path}")
    print(f"    Domains found: {found_count}/{row_count} rows")
    print(f"    WebLogic confirmed: {wls_count}/{row_count} rows")


def process_csv(input_csv: str, xlsx_path: str, row_range: tuple[int, int] | None = None):
    """Read CSV (IP:PORT,Domain), run httpx + nuclei, merge results into xlsx."""
    import csv

    # ── Read CSV ─────────────────────────────────────────────────────
    csv_data: dict[str, set[str]] = {}  # ip:port → {domains}
    ip_port_map: dict[str, tuple[str, int]] = {}
    with open(input_csv, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames or [])
        ip_col = None
        domain_col = None
        for fn in fieldnames:
            up = fn.upper().replace(" ", "")
            if "IP" in up and "PORT" in up:
                ip_col = fn
            elif "DOMAIN" in up or "HOST" in up or "FQDN" in up:
                domain_col = fn
        if not ip_col:
            ip_col = fieldnames[0]
        if not domain_col and len(fieldnames) > 1:
            domain_col = fieldnames[1]

        print(f"[*] CSV columns: IP={ip_col}, Domain={domain_col}")
        for row in reader:
            val = str(row.get(ip_col, "")).strip()
            if not val:
                continue
            if ":" in val:
                parts = val.rsplit(":", 1)
                ip, port = parts[0], int(parts[1])
            else:
                ip, port = val, 443
            ip_port_map[val] = (ip, port)
            domains: set[str] = set()
            if domain_col and row.get(domain_col):
                for d in str(row[domain_col]).split(","):
                    d = d.strip().lower()
                    if d and "." in d and not _is_ip(d):
                        domains.add(d)
            csv_data.setdefault(val, set()).update(domains)

    print(f"[*] Loaded {len(csv_data)} entries from {input_csv}")

    # ── httpx probe ──────────────────────────────────────────────────
    print("[*] Phase 1/2: httpx probe...")
    all_probe_targets: set[str] = set()
    for val, domains in csv_data.items():
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        for d in domains:
            all_probe_targets.add(f"{scheme}://{d}:{port}")
            all_probe_targets.add(f"{scheme}://{d}")
        all_probe_targets.add(f"{scheme}://{ip}:{port}")

    httpx_live: dict[str, dict] = {}
    if all_probe_targets:
        httpx_live = _run_httpx(sorted(all_probe_targets))
    print(f"    httpx: {len(httpx_live)} live targets")

    httpx_per_row: dict[str, str] = {}
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        domains = csv_data.get(val, set())
        live_domains = []
        for d in sorted(domains):
            for scheme in ("https", "http"):
                key = f"{scheme}://{d}:{port}"
                if key in httpx_live:
                    info = httpx_live[key]
                    live_domains.append(f"{d} [{info.get('status','')}]")
                    break
                key2 = f"{scheme}://{d}"
                if key2 in httpx_live:
                    info = httpx_live[key2]
                    live_domains.append(f"{d} [{info.get('status','')}]")
                    break
        if not live_domains:
            scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
            key = f"{scheme}://{ip}:{port}"
            if key in httpx_live:
                info = httpx_live[key]
                live_domains.append(f"{ip} [{info.get('status','')}]")
        httpx_per_row[val] = ", ".join(live_domains) if live_domains else ""

    # ── nuclei WLS fingerprint ───────────────────────────────────────
    print("[*] Phase 2/2: nuclei WebLogic fingerprint...")
    nuclei_targets: list[str] = []
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        nuclei_targets.append(f"{scheme}://{ip}:{port}")

    nuclei_results = _run_nuclei(sorted(set(nuclei_targets)))
    print(f"    nuclei: {len(nuclei_results)} confirmed WebLogic")

    nuclei_per_row: dict[str, str] = {}
    for val in ip_port_map:
        ip, port = ip_port_map[val]
        scheme = "https" if port in (443, 7002, 8443, 9002) else "http"
        key = f"{scheme}://{ip}:{port}"
        matches = nuclei_results.get(key, [])
        nuclei_per_row[val] = ", ".join(matches) if matches else ""

    # ── Merge into xlsx ──────────────────────────────────────────────
    print(f"[*] Merging results into {xlsx_path}...")
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active

    headers = [c.value for c in ws[1]]
    ip_col_idx = None
    for i, h in enumerate(headers):
        if h and "IP" in str(h).upper() and "PORT" in str(h).upper():
            ip_col_idx = i + 1
            break
    if ip_col_idx is None:
        print("ERROR: Could not find IP:PORT column in xlsx")
        sys.exit(1)

    # find or create the 3 result columns (Domain, httpx Probe, WLS Confirm)
    domain_ci = None
    httpx_ci = None
    wls_ci = None
    for i, h in enumerate(headers):
        if h == "Domain":
            domain_ci = i + 1
        elif h == "httpx Probe":
            httpx_ci = i + 1
        elif h == "WLS Confirm":
            wls_ci = i + 1

    if domain_ci is None:
        # columns don't exist yet — insert them after IP:PORT
        insert_at = ip_col_idx + 1
        ws.insert_cols(insert_at, amount=3)
        domain_ci = insert_at
        httpx_ci = insert_at + 1
        wls_ci = insert_at + 2
        src_hdr = ws.cell(row=1, column=ip_col_idx)
        for ci, name in [(domain_ci, "Domain"), (httpx_ci, "httpx Probe"), (wls_ci, "WLS Confirm")]:
            hdr = ws.cell(row=1, column=ci, value=name)
            copy_cell_style(src_hdr, hdr)
        for ci, w in [(domain_ci, 50), (httpx_ci, 40), (wls_ci, 35)]:
            ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w

    # merge: match by IP:PORT, update cells
    merge_min = 2
    merge_max = ws.max_row
    if row_range:
        merge_min = max(row_range[0], 2)
        merge_max = min(row_range[1], ws.max_row)
        print(f"[*] Merging into rows {merge_min}–{merge_max} only")

    updated = 0
    for row_idx in range(merge_min, merge_max + 1):
        cell_val = ws.cell(row=row_idx, column=ip_col_idx).value
        if not cell_val:
            continue
        val = str(cell_val).strip()
        if val not in csv_data:
            continue
        # Domain: append new domains to existing
        existing_domain = str(ws.cell(row=row_idx, column=domain_ci).value or "")
        existing_set = {d.strip().lower() for d in existing_domain.split(",") if d.strip()}
        merged = existing_set | csv_data[val]
        ws.cell(row=row_idx, column=domain_ci, value=", ".join(sorted(merged)) if merged else "")
        # httpx Probe: overwrite with fresh result
        ws.cell(row=row_idx, column=httpx_ci, value=httpx_per_row.get(val, ""))
        # WLS Confirm: overwrite with fresh result
        ws.cell(row=row_idx, column=wls_ci, value=nuclei_per_row.get(val, ""))
        updated += 1

    wb.save(xlsx_path)
    print(f"\n[+] Done! Updated {updated} rows in {xlsx_path}")
    print(f"    httpx live: {sum(1 for v in httpx_per_row.values() if v)}/{len(csv_data)}")
    print(f"    WebLogic confirmed: {sum(1 for v in nuclei_per_row.values() if v)}/{len(csv_data)}")


if __name__ == "__main__":
    import argparse

    XLSX = "Weblogic-VN-live-final.xlsx"

    parser = argparse.ArgumentParser(
        description="IP→Domain recon pipeline — writes results into " + XLSX)
    parser.add_argument("input", nargs="?", default=XLSX,
                        help="Input: xlsx (full recon) or csv (probe only)")
    parser.add_argument("--xlsx", default=XLSX,
                        help=f"Target xlsx to write results into (default: {XLSX})")
    parser.add_argument("--rows",
                        help="Row range to recon, e.g. 2-50 or 5-5 (header=row 1)")
    args = parser.parse_args()

    input_file = args.input
    xlsx_file = args.xlsx

    row_range = None
    if args.rows:
        parts = args.rows.split("-")
        if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
            row_range = (int(parts[0]), int(parts[1]))
        else:
            print(f"ERROR: invalid --rows format '{args.rows}', expected e.g. 2-50")
            sys.exit(1)

    if not Path(input_file).exists():
        print(f"ERROR: {input_file} not found")
        sys.exit(1)

    if input_file.endswith(".csv"):
        if not Path(xlsx_file).exists():
            print(f"ERROR: {xlsx_file} not found (needed to merge results into)")
            sys.exit(1)
        process_csv(input_file, xlsx_file, row_range=row_range)
    else:
        process(input_file, row_range=row_range)
