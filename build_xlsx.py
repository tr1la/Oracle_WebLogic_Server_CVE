#!/usr/bin/env python3
"""Join weblogic-confirmed-hostport.txt with WeblogicScan.json -> WebLogic-VN-confirmed.xlsx
Cột: URL (ưu tiên domain, xếp đầu), IP:PORT, Thành phố, Tổ chức.
Chạy: python3 build_xlsx.py   (cần: python3 -m pip install openpyxl --break-system-packages)
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))

def as_list(v):
    if isinstance(v, str): return [v] if v else []
    if isinstance(v, list): return v
    return []

def g(o, k):
    v = o.get(k); return v.strip() if isinstance(v, str) else ''

def host_domain(h):
    """Lấy phần tên miền từ 'host' của ZoomEye (vd 'apv1.dlp.com.tw:8888'); bỏ nếu là IP."""
    if not h: return ''
    name = h.rsplit(':', 1)[0].strip()
    # bỏ nếu là IPv4 thuần
    parts = name.split('.')
    if len(parts) == 4 and all(p.isdigit() for p in parts): return ''
    return name if ('.' in name and not name.replace('.', '').isdigit()) else ''

def normalize(o):
    """Chuẩn hóa 1 bản ghi FOFA hoặc ZoomEye về cùng schema."""
    port = str(o.get('port') or '')
    if 'city_name_EN' in o or 'organization' in o:          # --- FOFA ---
        site = g(o, 'site'); rdns = g(o, 'rdns'); doms = as_list(o.get('domains'))
        dom = site or rdns or (str(doms[0]).strip() if doms else '')   # FOFA: ưu tiên site -> rdns -> domains
        rec = {'domain': dom, 'rdns': rdns, 'city': g(o, 'city_name_EN'),
               'org': g(o, 'organization'), 'service': g(o, 'service')}
    else:                                                    # --- ZoomEye ---
        dom = g(o, 'domain') or host_domain(g(o, 'host'))
        rec = {'domain': dom, 'rdns': '', 'city': g(o, 'city'),
               'org': g(o, 'org'), 'service': g(o, 'protocol')}
    return port, rec

agg = {}; ports_of = {}
with open(os.path.join(BASE, "WeblogicScan.json")) as f:
    for line in f:
        line = line.strip()
        if not line: continue
        o = json.loads(line)
        port, rec = normalize(o)
        for ip in as_list(o.get('ip')):
            key = (ip, port); cur = agg.get(key)
            if cur is None:
                agg[key] = dict(rec)
            else:
                for k, v in rec.items():
                    if v and not cur.get(k): cur[k] = v
            ports_of.setdefault(ip, set()).add(port)

FIELDS = ('domain', 'rdns', 'city', 'org', 'service')

def lookup(ip, port):
    merged = {k: '' for k in FIELDS}
    if port and (ip, port) in agg:                       # 1) bản ghi khớp đúng ip:port
        for k in FIELDS:
            if agg[(ip, port)].get(k): merged[k] = agg[(ip, port)][k]
    order = [p for p in ('443', '80') if p in ports_of.get(ip, set())] + \
            [p for p in ports_of.get(ip, set()) if p not in ('443', '80')]
    for p in order:                                      # 2) lấp field rỗng từ cổng khác cùng IP
        for k in FIELDS:
            if not merged[k] and agg[(ip, p)].get(k): merged[k] = agg[(ip, p)][k]
    return merged

rows = []
with open(os.path.join(BASE, "weblogic-confirmed-hostport.txt")) as f:
    for line in f:
        hp = line.strip()
        if not hp: continue
        ip, port = (hp.rsplit(':', 1) if ':' in hp else (hp, ''))
        r = lookup(ip, port)
        svc = r['service'].lower()
        https = (port in ('443', '8443', '7002')) or ('https' in svc) or ('ssl' in svc)
        scheme = 'https' if https else 'http'
        host = r['domain'] if r['domain'] else ip
        pp = '' if port in ('', '80', '443') else ':' + port
        rows.append({'url': f"{scheme}://{host}{pp}", 'has_domain': bool(r['domain']),
                     'ipport': hp if port else ip, 'city': r['city'], 'org': r['org']})

rows.sort(key=lambda x: (not x['has_domain'], x['url']))   # domain lên đầu

wb = Workbook(); ws = wb.active; ws.title = "WebLogic VN"
ws.append(["URL", "IP:PORT", "Thành phố", "Tổ chức"])
hf = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="C0392B")
thin = Side(style="thin", color="DDDDDD"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
for c in ws[1]:
    c.font = hf; c.fill = fill; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = border
dom_fill = PatternFill("solid", fgColor="FDF2E9")
for r in rows:
    ws.append([r['url'], r['ipport'], r['city'], r['org']])
for i, r in enumerate(rows, start=2):
    for c in ws[i]:
        c.border = border; c.alignment = Alignment(vertical="center")
        if r['has_domain']: c.fill = dom_fill
for j, w in enumerate([46, 24, 18, 48], start=1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:D{len(rows)+1}"
out = os.path.join(BASE, "WebLogic-VN-confirmed.xlsx")
wb.save(out)

ndom = sum(1 for r in rows if r['has_domain'])
nocity = sum(1 for r in rows if not r['city'])
print(f"Da ghi {out}")
print(f"Tong: {len(rows)} | co domain: {ndom} | thieu city: {nocity}")
